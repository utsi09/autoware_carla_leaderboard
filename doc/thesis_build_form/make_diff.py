# -*- coding: utf-8 -*-
"""수정 표시판(diff): 지울 부분=빨강, 추가할 부분=초록(바로 옆/아래 삽입).
원본(사용자 0831 docx)은 구조 그대로 유지."""
import os, copy
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import Table

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(HERE, 'user_submit_base.docx')
OUT = '/home/taewook/autoware_carla_leaderboard/doc/thesis_diff_v1_0901.docx'
FONT = 'HY신명조'
RED = RGBColor(0xC0, 0x00, 0x00)
GREEN = RGBColor(0x00, 0x8A, 0x00)

doc = Document(BASE)

def find_para(text, start=0, exact=False):
    for i, p in enumerate(doc.paragraphs):
        if i >= start and ((p.text.strip() == text) if exact else (text in p.text)):
            return i, p
    raise KeyError(text)

def color_el(el, rgb):
    """요소(문단/표) 내 모든 런에 색 적용"""
    for r in el.iter(qn('w:r')):
        rPr = r.find(qn('w:rPr'))
        if rPr is None:
            rPr = OxmlElement('w:rPr'); r.insert(0, rPr)
        c = rPr.find(qn('w:color'))
        if c is None:
            c = OxmlElement('w:color'); rPr.append(c)
        c.set(qn('w:val'), '%02X%02X%02X' % (rgb[0], rgb[1], rgb[2]))

def green_run(p, text, size=11, bold=False):
    r = p.add_run(text)
    r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = GREEN
    rPr = r._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rPr.insert(0, rf)
    rf.set(qn('w:eastAsia'), FONT)
    return r

# ---------- 1) 유의사항 페이지 전체 빨강 ----------
i, p = find_para('학위논문 작성 유의사항')
j, stop = find_para('본 페이지는 전자전기컴퓨터공학과')
el = p._element
while el is not None and el is not stop._element:
    color_el(el, RED)
    el = el.getnext()
print('1) 유의사항 빨강')

# ---------- 2) 측면표지 빨강 ----------
i, sc = find_para('Side Cover')
el = sc._element
while el is not None:
    if not el.tag.endswith('}sectPr'):
        color_el(el, RED)
    el = el.getnext()
print('2) 측면표지 빨강')

# ---------- 3) 6장 절번호: 빨강 + 초록 대체 ----------
i, p = find_para('6. Limitations', exact=True)
color_el(p._p, RED)
green_run(p, '   →   4. Limitations', size=14, bold=True)
print('3) 6.→4. Limitations')

# ---------- 4) 4장 헤딩 2건 ----------
for old, new in [('4.1 Experimental Environment', '1. Experimental Environment'),
                 ('4.2 Fault Model', '2. Fault Model')]:
    i, p = find_para(old, exact=True)
    color_el(p._p, RED)
    green_run(p, f'   →   {new}', size=14, bold=True)
print('4) 4장 헤딩')

# ---------- 5) 2장 도입문 ----------
i, p = find_para('six strands (Sections 2.1-2.6)')
color_el(p._p, RED)
green_run(p, '  →  Related work is organized into four strands (Sections 2.1-2.4), '
             'and Table 2.1 compares them with this study.', size=11)
print('5) 2장 도입문')

# ---------- 6) TOC/LoT/LoF: 기존 표 빨강 + 새 표(초록) 바로 아래 삽입 ----------
TOC_GRID = [2450, 849, 426, 1134, 993, 1276, 809, 567]
LO_GRID = [2127, 5811, 567]
DOTS = {7: 28, 6: 24, 5: 20, 4: 18, 3: 13, 2: 9, 1: 4}

def _mk_tc(width, span, text, jc, green=True):
    tc = OxmlElement('w:tc')
    tcPr = OxmlElement('w:tcPr')
    w = OxmlElement('w:tcW'); w.set(qn('w:w'), str(width)); w.set(qn('w:type'), 'dxa'); tcPr.append(w)
    if span > 1:
        gs = OxmlElement('w:gridSpan'); gs.set(qn('w:val'), str(span)); tcPr.append(gs)
    tb = OxmlElement('w:tcBorders')
    for s in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{s}'); e.set(qn('w:val'), 'nil'); tb.append(e)
    tcPr.append(tb)
    tm = OxmlElement('w:tcMar')
    for s, v in (('top', 28), ('left', 102), ('bottom', 28), ('right', 102)):
        e = OxmlElement(f'w:{s}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); tm.append(e)
    tcPr.append(tm)
    va = OxmlElement('w:vAlign'); va.set(qn('w:val'), 'center'); tcPr.append(va)
    tc.append(tcPr)
    pel = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    pPr.append(OxmlElement('w:wordWrap'))
    sg = OxmlElement('w:snapToGrid'); sg.set(qn('w:val'), '0'); pPr.append(sg)
    sp = OxmlElement('w:spacing'); sp.set(qn('w:line'), '348'); sp.set(qn('w:lineRule'), 'auto'); pPr.append(sp)
    if jc:
        jc_ = OxmlElement('w:jc'); jc_.set(qn('w:val'), jc); pPr.append(jc_)
    pel.append(pPr)
    rel = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts'); rf.set(qn('w:ascii'), FONT); rf.set(qn('w:eastAsia'), FONT); rPr.append(rf)
    if green:
        c = OxmlElement('w:color'); c.set(qn('w:val'), '008A00'); rPr.append(c)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '22'); rPr.append(sz)
    rel.append(rPr)
    t_ = OxmlElement('w:t'); t_.set(qn('xml:space'), 'preserve'); t_.text = text; rel.append(t_)
    pel.append(rel)
    tc.append(pel)
    return tc

def new_tbl_el(grid, rows3):
    te = OxmlElement('w:tbl')
    tblPr = OxmlElement('w:tblPr')
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), str(sum(grid))); w.set(qn('w:type'), 'dxa'); tblPr.append(w)
    bd = OxmlElement('w:tblBorders')
    for s in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{s}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '2')
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), '008A00'); bd.append(e)
    for s in ('insideH', 'insideV'):
        e = OxmlElement(f'w:{s}'); e.set(qn('w:val'), 'nil'); bd.append(e)
    tblPr.append(bd)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    te.append(tblPr)
    g = OxmlElement('w:tblGrid')
    for cw in grid:
        c = OxmlElement('w:gridCol'); c.set(qn('w:w'), str(cw)); g.append(c)
    te.append(g)
    for cells in rows3:
        tr = OxmlElement('w:tr')
        for tc in cells:
            tr.append(tc)
        te.append(tr)
    return te

def toc_rows(entries):
    cum = [sum(TOC_GRID[:k]) for k in range(1, 8)]
    rows = []
    for text, page in entries:
        est = int(len(text) * 128) + 204
        k = next((i + 1 for i, c in enumerate(cum[:6]) if c >= est), 6)
        rows.append([_mk_tc(cum[k - 1], k, text, 'left'),
                     _mk_tc(sum(TOC_GRID[k:7]), 7 - k, '·' * DOTS[7 - k], 'center'),
                     _mk_tc(567, 1, str(page), 'center')])
    return rows

def lo_rows(entries):
    return [[_mk_tc(LO_GRID[0], 1, label, 'left'),
             _mk_tc(LO_GRID[1], 1, '·' * 25, 'center'),
             _mk_tc(567, 1, str(page), 'center')] for label, page in entries]

def table_after(text):
    i, p = find_para(text)
    el = p._element.getnext()
    while el is not None and not el.tag.endswith('}tbl'):
        el = el.getnext()
    return el

toc = [
 ('Abstract', 'v'), ('1. Introduction', '1'), ('1.1  Research Background', '1'),
 ('1.2  Research Questions', '3'), ('1.3  Contribution', '4'),
 ('2. Background and Related Work', '5'), ('2.1  Standards and Regulatory Frameworks', '6'),
 ('2.2  FTTI and the Derivation of Time Budgets', '6'),
 ('2.3  Fault Detection Theory and Sensor Degradation Models', '7'),
 ('2.4  Fallback and MRM Planning', '7'),
 ('3. Proposed System', '11'), ('3.1  System Architecture', '11'), ('3.2  Sensor Scoring', '13'),
 ('3.3  Spatial Sectors and Capability Synthesis', '14'),
 ('3.4  Risk-Coupled Dynamic Detection Budget', '15'), ('3.5  MRM Supervisor', '17'),
 ('3.6  Conservative Low-Speed MRM Movement (Speed Cap)', '21'),
 ('3.7  Implementation Boundary and Map Annotation', '22'),
 ('4. Experimental Environment and Design', '22'), ('4.1  Experimental Environment', '23'),
 ('4.2  Fault Model', '24'), ('4.3  Main Campaign Matrix', '27'),
 ('4.4  Pedestrian Dart-out Encounter Injector', '28'),
 ('4.5  Collision Adjudication Methodology', '28'),
 ('4.6  Distance Ladder and 8 m Encounter Experiment Design', '29'),
 ('4.7  Metrics and Statistics', '30'), ('4.8  Reproducibility: Key Experimental Parameters', '31'),
 ('5. Experimental Results', '32'),
 ('5.1  Main Campaign: Detection Performance and Per-Fault Behavior', '32'),
 ('5.2  Exposure Metric: Degraded-State Driving Distance', '38'),
 ('5.3  Distance Ladder: Measuring the Avoid/Collide Boundary', '38'),
 ('5.4  The 8 m Encounter: Low-Speed MRM versus Stock', '40'),
 ('5.5  Low-Speed Ladder: Shift of the Boundary', '42'),
 ('5.6  Causal Verification of the Speed Cap: Dose-Response Experiment', '42'),
 ('5.7  Hard-Fail Decision Latency: Verifying the 1.5 s Deadline', '45'),
 ('6. Limitations and Discussion', '46'),
 ('6.1  Temporal Trade-off, Transient Risk, and Speed Cap', '46'),
 ('6.2  Lessons for Adjudication Methodology', '47'),
 ('6.3  The Cost of Design Conservatism and the Execution/Map Layers', '47'),
 ('6.4  Limitations', '48'), ('7. Conclusion', '49'), ('References', '50'),
 ('논문요약 (Korean Abstract)', '57'),
]
lot = [('Table 2.1.', '9'), ('Table 4.1.', '25'), ('Table 4.2.', '27'), ('Table 4.3.', '31'),
       ('Table 5.1.', '33'), ('Table 5.2.', '37'), ('Table 5.3.', '39'), ('Table 5.4.', '43'), ('Table 5.5.', '45')]
lof = [('Fig. 1.1.', '3'), ('Fig. 3.1.', '12'), ('Fig. 3.2.', '15'), ('Fig. 3.3.', '18'), ('Fig. 3.4.', '20'),
       ('Fig. 4.1.', '23'), ('Fig. 4.2.', '26'), ('Fig. 5.1.', '35'), ('Fig. 5.2.', '37'), ('Fig. 5.3.', '41'), ('Fig. 5.4.', '44')]

for head, grid, rows in [('Table of Contents', TOC_GRID, toc_rows(toc)),
                         ('List of Tables', LO_GRID, lo_rows(lot)),
                         ('List of Figures', LO_GRID, lo_rows(lof))]:
    old = table_after(head)
    color_el(old, RED)
    new = new_tbl_el(grid, rows)
    old.addnext(new)
print('6) 목차 3종: 기존 빨강 + 새 표 초록')

doc.save(OUT)
print('saved', OUT)
