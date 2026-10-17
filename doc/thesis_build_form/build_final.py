# -*- coding: utf-8 -*-
"""사용자 최종 작업본(user3) 기반 완성 빌더: 표 4.2부터 끝까지 채움.
- 사용자가 손본 부분(~표 4.2 캡션, 그림 크기, 수식 (1)~(9))은 무수정 보존
- 수정: Eq. (3.1) 잔재 → Eq. (1)
- 추가: 표 4.2 본체+주석, 4.3 꼬리 문단, 4.4~4.8, Chapter 5~7(대문자 제목),
  References, 국문 논문요약, TOC/LoT/LoF 갱신(2-pass 쪽수)
"""
import json, os, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import Table

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from content_en1 import SUMMARY_KO_TITLE, SUMMARY_KO, SUMMARY_KO_KEYWORDS
from content_en2 import CH4
from content_en3 import CH5, CH6, CH7, REFS

BASE = os.path.join(HERE, 'user_final_base.docx')
OUT = '/home/taewook/autoware_carla_leaderboard/doc/thesis_bonsim_v3_0831.docx'
FONT = 'HY신명조'

pagemap = {}
pm = os.path.join(HERE, 'pagemap_final.json')
if os.path.exists(pm):
    pagemap = json.load(open(pm))
keys = {}

KO_TITLE_L = ['자율주행 차량의 페일세이프 운행을 위한', '연속 센서 건강도 감시와',
              '위험·잔여능력 기반 단계적 최소위험기동 감독']

doc = Document(BASE)

def styled(r, size=11, bold=False):
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    rPr = r._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rPr.insert(0, rf)
    rf.set(qn('w:eastAsia'), FONT)
    return r

def find_para(text, start=0):
    for i, p in enumerate(doc.paragraphs):
        if i >= start and text in p.text:
            return i, p
    raise KeyError(text)

def new_para(text='', size=11, bold=False, align=None, ls=1.45, fi=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = ls
    pf.space_before = Pt(0); pf.space_after = Pt(0)
    if align is not None: p.alignment = align
    if fi is not None: pf.first_line_indent = Pt(fi)
    if text:
        styled(p.add_run(text), size=size, bold=bold)
    return p

def new_figure(caption, key, path, width_cm):
    out = []
    if path and os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Cm(width_cm or 14.5))
        out.append(p)
    cap = new_para(caption, size=11, bold=True, ls=1.3)
    cap.paragraph_format.keep_together = True
    out.append(cap)
    out.append(new_para(''))
    keys[key] = caption.split('. ')[0] + '.'
    return out

def new_table(caption, key, rows, note, with_caption=True):
    out = []
    if with_caption:
        cp = new_para(caption, size=11, bold=True, ls=1.3)
        cp.paragraph_format.keep_with_next = True
        out.append(cp)
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = 1
    for ri, tr in enumerate(t.rows):
        trPr = tr._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        if ri == 0:
            trPr.append(OxmlElement('w:tblHeader'))
    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            cell = t.cell(ri, ci); cell.text = ''
            first = True
            for line in str(cell_text).split('\n'):
                cp2 = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                cp2.paragraph_format.line_spacing = 1.2
                cp2.paragraph_format.space_before = Pt(1)
                cp2.paragraph_format.space_after = Pt(1)
                cp2.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
                styled(cp2.add_run(line), size=9.0, bold=(ri == 0))
    out.append(t)
    if note:
        out.append(new_para(note, size=9.0, ls=1.3))
    out.append(new_para(''))
    keys[key] = caption.split('. ')[0] + '.'
    return out

def render_blocks(blocks):
    out = []
    for b in blocks:
        kind = b[0]
        if kind == 'ch':
            head, tail = b[1].split('. ', 1)
            title = head + '. ' + tail.upper()
            keys[b[2]] = title
            p = new_para('', ls=1.45)
            r = p.add_run(); r.add_break(WD_BREAK.PAGE)
            styled(r, 16, True)
            styled(p.add_run(title), 16, True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            out.append(p); out.append(new_para(''))
        elif kind == 'sec':
            keys[b[2]] = b[1]
            out.append(new_para(''))
            sp = new_para(b[1], size=14, bold=True)
            sp.paragraph_format.keep_with_next = True
            out.append(sp); out.append(new_para(''))
        elif kind == 'sub':
            sp = new_para(b[1], size=11, fi=11.0)
            sp.paragraph_format.keep_with_next = True
            out.append(sp)
        elif kind == 'p':
            pp = new_para(b[1], size=11, fi=16.5)
            pp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            out.append(pp)
        elif kind == 'fig':
            out.extend(new_figure(b[1], b[2], b[3], b[5] if len(b) > 5 else None))
        elif kind == 'tbl':
            out.extend(new_table(b[1], b[2], b[3], b[4]))
    return out

def move_before(anchor_el, elements):
    for el in elements:
        node = el._element if hasattr(el, '_element') else el._p
        anchor_el.addprevious(node)

# ---------- 목차 표 (v2와 동일 구조) ----------
def _mk_tc(width, span, text, jc, bold=False):
    tc = OxmlElement('w:tc')
    tcPr = OxmlElement('w:tcPr')
    w = OxmlElement('w:tcW'); w.set(qn('w:w'), str(width)); w.set(qn('w:type'), 'dxa'); tcPr.append(w)
    if span > 1:
        gs = OxmlElement('w:gridSpan'); gs.set(qn('w:val'), str(span)); tcPr.append(gs)
    tb = OxmlElement('w:tcBorders')
    for s in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{s}'); e.set(qn('w:val'), 'nil'); tb.append(e)
    tcPr.append(tb)
    tm = OxmlElement('w:tcMar')
    for s, v in (('top', 28), ('left', 102), ('bottom', 28), ('right', 102)):
        e = OxmlElement(f'w:{s}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); tm.append(e)
    tcPr.append(tm)
    va = OxmlElement('w:vAlign'); va.set(qn('w:val'), 'center'); tcPr.append(va)
    tc.append(tcPr)
    pel = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    pPr.append(OxmlElement('w:wordWrap'))
    sg = OxmlElement('w:snapToGrid'); sg.set(qn('w:val'), '0'); pPr.append(sg)
    sp = OxmlElement('w:spacing'); sp.set(qn('w:line'), '348'); sp.set(qn('w:lineRule'), 'auto'); pPr.append(sp)
    if jc:
        j = OxmlElement('w:jc'); j.set(qn('w:val'), jc); pPr.append(j)
    pel.append(pPr)
    rel = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts'); rf.set(qn('w:ascii'), FONT); rf.set(qn('w:eastAsia'), FONT); rPr.append(rf)
    if bold:
        rPr.append(OxmlElement('w:b'))
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '22'); rPr.append(sz)
    rel.append(rPr)
    t_ = OxmlElement('w:t'); t_.set(qn('xml:space'), 'preserve'); t_.text = text; rel.append(t_)
    pel.append(rel)
    tc.append(pel)
    return tc

TOC_GRID = [2450, 849, 426, 1134, 993, 1276, 809, 567]
LO_GRID = [2127, 5811, 567]
DOTS = {7: 28, 6: 24, 5: 20, 4: 18, 3: 13, 2: 9, 1: 4}

def _set_grid_and_borders(te, grid):
    tblPr = te.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr'); te.insert(0, tblPr)
    for tag in ('w:tblW', 'w:tblBorders', 'w:tblLayout', 'w:tblCellMar'):
        e = tblPr.find(qn(tag))
        if e is not None:
            tblPr.remove(e)
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), str(sum(grid))); w.set(qn('w:type'), 'dxa'); tblPr.append(w)
    bd = OxmlElement('w:tblBorders')
    for s in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{s}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '2')
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), '000000'); bd.append(e)
    for s in ('insideH', 'insideV'):
        e = OxmlElement(f'w:{s}'); e.set(qn('w:val'), 'nil'); bd.append(e)
    tblPr.append(bd)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    cm = OxmlElement('w:tblCellMar')
    for s, v in (('top', 15), ('left', 15), ('bottom', 15), ('right', 15)):
        e = OxmlElement(f'w:{s}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); cm.append(e)
    tblPr.append(cm)
    og = te.find(qn('w:tblGrid'))
    if og is not None:
        te.remove(og)
    g = OxmlElement('w:tblGrid')
    for cw in grid:
        c = OxmlElement('w:gridCol'); c.set(qn('w:w'), str(cw)); g.append(c)
    te.insert(list(te).index(tblPr) + 1, g)

def fill_toc_table(tbl, entries):
    te = tbl._tbl
    _set_grid_and_borders(te, TOC_GRID)
    for tr in list(te.findall(qn('w:tr'))):
        te.remove(tr)
    cum = [sum(TOC_GRID[:k]) for k in range(1, 8)]
    for text, page in entries:
        est = int(len(text) * 128) + 204
        k = next((i + 1 for i, c in enumerate(cum[:6]) if c >= est), 6)
        dots_span = 7 - k
        tr = OxmlElement('w:tr')
        tr.append(_mk_tc(cum[k - 1], k, text, 'left'))
        tr.append(_mk_tc(sum(TOC_GRID[k:7]), dots_span, '·' * DOTS[dots_span], 'center'))
        tr.append(_mk_tc(567, 1, str(page), 'center'))
        te.append(tr)

def fill_lo_table(tbl, entries):
    te = tbl._tbl
    _set_grid_and_borders(te, LO_GRID)
    for tr in list(te.findall(qn('w:tr'))):
        te.remove(tr)
    for label, page in entries:
        tr = OxmlElement('w:tr')
        tr.append(_mk_tc(LO_GRID[0], 1, label, 'left'))
        tr.append(_mk_tc(LO_GRID[1], 1, '·' * 25, 'center'))
        tr.append(_mk_tc(567, 1, str(page), 'center'))
        te.append(tr)

def table_after_para(text):
    i, p = find_para(text)
    el = p._element.getnext()
    while el is not None and not el.tag.endswith('}tbl'):
        el = el.getnext()
    return Table(el, doc)

# ================= 1) Eq. (3.1) 잔재 수정 =================
i, p = find_para('decay constant of Eq. (3.1)')
for t_el in p._p.iter(qn('w:t')):
    if t_el.text and '(3.1)' in t_el.text:
        t_el.text = t_el.text.replace('(3.1)', '(1)')
print('fixed Eq.(3.1) -> Eq.(1)')

# ================= 1.5) 장 제목 새 페이지 시작 (사용자 붙임분 포함) =================
_HEADS = ['Abstract', 'Chapter 1. INTRODUCTION', 'Chapter 2. BACKGROUND AND RELATED WORK',
          'Chapter 3. PROPOSED SYSTEM', 'Chapter 4. EXPERIMENTAL ENVIRONMENT AND DESIGN']
for _h in _HEADS:
    for _p in doc.paragraphs:
        if _p.text.strip() == _h:
            _p.paragraph_format.page_break_before = True
            break

# ================= 2) 표 4.2 본체 + 4장 나머지 + 5~7장 =================
i, cap = find_para('Table 4.2. Main-campaign case accounting')
iR, pR = find_para('References (16p)')
# CH4에서 tbl4_acct부터 끝까지 슬라이스
idx = next(k for k, b in enumerate(CH4) if b[0] == 'tbl' and b[2] == 'tbl4_acct')
tail = CH4[idx:]
# 표 4.2는 캡션이 이미 있으므로 캡션 없이
gen = []
b = tail[0]
gen.extend(new_table(b[1], b[2], b[3], b[4], with_caption=False))
gen.extend(render_blocks(tail[1:]))
gen.extend(render_blocks(CH5))
gen.extend(render_blocks(CH6))
gen.extend(render_blocks(CH7))
move_before(pR._element, gen)
# 캡션과 생성 콘텐츠 사이의 양식 샘플 잔재 삭제
first_gen = gen[0]._element if hasattr(gen[0], '_element') else gen[0]._tbl
el = cap._element.getnext()
while el is not None and el is not first_gen:
    nxt = el.getnext()
    el.getparent().remove(el)
    el = nxt
keys['ch4s3t'] = 'Table 4.2.'

# ================= 3) References =================
pR.text = ''
pR.alignment = WD_ALIGN_PARAGRAPH.CENTER
pR.paragraph_format.line_spacing = 1.45
r = pR.add_run(); r.add_break(WD_BREAK.PAGE); styled(r, 16, True)
styled(pR.add_run('References'), 16, True)
keys['refs'] = 'References'
iA, pA = find_para('Appendix (16p)')
el = pR._element.getnext()
while el is not None and el is not pA._element:
    nxt = el.getnext()
    el.getparent().remove(el)
    el = nxt
ref_paras = []
for n, ref in enumerate(REFS, 1):
    rp = new_para(f'{n}. {ref}', size=11, ls=1.3)
    rp.paragraph_format.left_indent = Cm(0.9)
    rp.paragraph_format.first_line_indent = Cm(-0.9)
    rp.paragraph_format.space_after = Pt(3)
    ref_paras.append(rp)
move_before(pA._element, ref_paras)

# ================= 4) Appendix 삭제 + 논문요약 =================
found = None
for pp in doc.paragraphs:
    if pp.text.strip() == '논문요약':
        found = pp
assert found is not None
brk = new_para('')
_r = brk.add_run(); _r.add_break(WD_BREAK.PAGE)
found._element.addprevious(brk._p)
el = pA._element
while el is not None and el is not found._element:
    nxt = el.getnext()
    if el is brk._p:
        el = nxt; continue
    el.getparent().remove(el)
    el = nxt
keys['summary_ko'] = '논문요약'

i, pT = find_para('논문제목을 작성하시오')
pT.text = ''
pT.alignment = WD_ALIGN_PARAGRAPH.CENTER
pT.paragraph_format.line_spacing = 1.45
for li, line in enumerate(KO_TITLE_L):
    r = styled(pT.add_run(line), 16, True)
    if li < 2:
        r.add_break()
for anchor, val in [('제출자 성명', '엄태욱'), ('OO학과', '전자전기컴퓨터공학과'), ('성균관대학교(11p)', '성균관대학교')]:
    i, pp = find_para(anchor)
    pp.text = ''
    pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pp.paragraph_format.line_spacing = 1.45
    styled(pp.add_run(val), 11)
i0, p0 = find_para('여기부터 본문을 국문으로')
iK, pK = find_para('주제어:', start=i0)
ko_paras = [new_para(t, size=11, fi=16.5, ls=1.45) for t in SUMMARY_KO]
for kp in ko_paras:
    kp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
move_before(p0._element, ko_paras)
el = p0._element
while el is not None and el is not pK._element:
    nxt = el.getnext()
    if el.tag.endswith('}p'):
        el.getparent().remove(el)
    el = nxt
pK.text = ''
styled(pK.add_run(SUMMARY_KO_KEYWORDS), 11, True)
pK.paragraph_format.line_spacing = 1.45

# ================= 5) 기존(사용자 작성) 헤딩·캡션 키 등록 =================
keys.update({
 'abstract': 'Fail-Safe Operation of Autonomous Vehicles',
 'ch1': 'Chapter 1. INTRODUCTION', 'ch1s1': '1.1 Research Background',
 'ch1s2': '1.2 Research Questions', 'ch1s3': '1.3 Contribution',
 'ch1s4': '1.4 Thesis Organization',
 'ch2': 'Chapter 2. BACKGROUND AND RELATED WORK',
 'ch2s1': '2.1 Standards and Regulatory Frameworks',
 'ch2s2': '2.2 FTTI and the Derivation of Time Budgets',
 'ch2s3': '2.3 Fault Detection Theory and Sensor Degradation Models',
 'ch2s4': '2.4 Fallback and MRM Planning',
 'ch2s5': '2.5 Capability-Based Self-Monitoring and Runtime Monitoring',
 'ch2s6': '2.6 Simulation-Based Evaluation and Fault Injection',
 'ch2s7': '2.7 Comparison with Prior Work',
 'ch3': 'Chapter 3. PROPOSED SYSTEM', 'ch3s1': '3.1 System Architecture',
 'ch3s2': '3.2 Sensor Scoring', 'ch3s3': '3.3 Spatial Sectors and Capability Synthesis',
 'ch3s4': '3.4 Risk-Coupled Dynamic Detection Budget', 'ch3s5': '3.5 MRM Supervisor',
 'ch3s6': '3.6 Remaining-Capability-Based Conservative Low-Speed MRM Movement',
 'ch3s7': '3.7 MRM Candidate Map Annotation',
 'ch3s8': '3.8 Boundary and Interface with the Execution Layer',
 'ch4': 'Chapter 4. EXPERIMENTAL ENVIRONMENT AND DESIGN',
 'ch4s1': '4.1 Experimental Environment', 'ch4s2': '4.2 Fault Model',
 'ch4s3': '4.3 Main Campaign Matrix',
 'tbl2_1': 'Table 2.1.', 'tbl4_1': 'Table 4.1.',
 'fig1_1': 'Figure 1.1.', 'fig3_1': 'Figure 3.1.', 'fig3_bt': 'Figure 3.2.',
 'fig3_2': 'Figure 3.3.', 'fig3_4': 'Figure 3.4.', 'fig4_2': 'Figure 4.1.',
 'fig4_faults': 'Figure 4.2.',
})

# ================= 6) TOC / LoT / LoF =================
def pg(key, default='0'):
    return pagemap.get(key, default)

toc_entries = [
    ('Abstract', pg('abstract', 'v')),
    ('1. Introduction', pg('ch1', '1')),
    ('1.1  Research Background', pg('ch1s1', '1')),
    ('1.2  Research Questions', pg('ch1s2', '3')),
    ('1.3  Contribution', pg('ch1s3', '4')),
    ('1.4  Thesis Organization', pg('ch1s4', '5')),
    ('2. Background and Related Work', pg('ch2', '6')),
    ('2.1  Standards and Regulatory Frameworks', pg('ch2s1', '6')),
    ('2.2  FTTI and the Derivation of Time Budgets', pg('ch2s2', '6')),
    ('2.3  Fault Detection Theory and Sensor Degradation Models', pg('ch2s3', '7')),
    ('2.4  Fallback and MRM Planning', pg('ch2s4', '7')),
    ('2.5  Capability-Based Self-Monitoring and Runtime Monitoring', pg('ch2s5', '9')),
    ('2.6  Simulation-Based Evaluation and Fault Injection', pg('ch2s6', '9')),
    ('2.7  Comparison with Prior Work', pg('ch2s7', '9')),
    ('3. Proposed System', pg('ch3', '12')),
    ('3.1  System Architecture', pg('ch3s1', '12')),
    ('3.2  Sensor Scoring', pg('ch3s2', '13')),
    ('3.3  Spatial Sectors and Capability Synthesis', pg('ch3s3', '15')),
    ('3.4  Risk-Coupled Dynamic Detection Budget', pg('ch3s4', '15')),
    ('3.5  MRM Supervisor', pg('ch3s5', '17')),
    ('3.6  Conservative Low-Speed MRM Movement (Speed Cap)', pg('ch3s6', '20')),
    ('3.7  MRM Candidate Map Annotation', pg('ch3s7', '21')),
    ('3.8  Boundary and Interface with the Execution Layer', pg('ch3s8', '21')),
    ('4. Experimental Environment and Design', pg('ch4', '23')),
    ('4.1  Experimental Environment', pg('ch4s1', '23')),
    ('4.2  Fault Model', pg('ch4s2', '24')),
    ('4.3  Main Campaign Matrix', pg('ch4s3', '26')),
    ('4.4  Pedestrian Dart-out Encounter Injector', pg('ch4s4', '27')),
    ('4.5  Collision Adjudication Methodology', pg('ch4s5', '28')),
    ('4.6  Distance Ladder and 8 m Encounter Experiment Design', pg('ch4s6', '29')),
    ('4.7  Metrics and Statistics', pg('ch4s7', '30')),
    ('4.8  Reproducibility: Key Experimental Parameters', pg('ch4s8', '30')),
    ('5. Experimental Results', pg('ch5', '32')),
    ('5.1  Main Campaign: Detection Performance and Per-Fault Behavior', pg('ch5s1', '32')),
    ('5.2  Exposure Metric: Degraded-State Driving Distance', pg('ch5s2', '36')),
    ('5.3  Distance Ladder: Measuring the Avoid/Collide Boundary', pg('ch5s3', '37')),
    ('5.4  The 8 m Encounter: Low-Speed MRM versus Stock', pg('ch5s4', '39')),
    ('5.5  Low-Speed Ladder: Shift of the Boundary', pg('ch5s5', '40')),
    ('5.6  Causal Verification of the Speed Cap: Dose-Response Experiment', pg('ch5s6', '40')),
    ('5.7  Hard-Fail Decision Latency: Verifying the 1.5 s Deadline', pg('ch5s7', '43')),
    ('6. Limitations and Discussion', pg('ch6', '45')),
    ('6.1  The Temporal Trade-off Between Stopping and Evacuation', pg('ch6s1', '45')),
    ('6.2  The Post-Detection Transient, the Risk Window, and FTTI', pg('ch6s2', '45')),
    ('6.3  The Role of the Speed Cap: Causal Effect and Limits', pg('ch6s3', '46')),
    ('6.4  Lessons for Adjudication Methodology', pg('ch6s4', '46')),
    ('6.5  The Cost of Design Conservatism and the Execution/Map Layers', pg('ch6s5', '47')),
    ('6.6  Limitations', pg('ch6s6', '47')),
    ('7. Conclusion', pg('ch7', '49')),
    ('References', pg('refs', '51')),
    ('논문요약 (Korean Abstract)', pg('summary_ko', '56')),
]
lot_entries = [
    ('Table 2.1.', pg('tbl2_1', '10')),
    ('Table 4.1.', pg('tbl4_1', '25')),
    ('Table 4.2.', pg('ch4s3t', '27')),
    ('Table 4.3.', pg('tbl4_3', '31')),
    ('Table 5.1.', pg('tbl5_0', '34')),
    ('Table 5.2.', pg('tbl5_1', '37')),
    ('Table 5.3.', pg('tbl5_2', '39')),
    ('Table 5.4.', pg('tbl5_sd', '42')),
    ('Table 5.5.', pg('tbl5_7', '44')),
]
lof_entries = [
    ('Figure 1.1.', pg('fig1_1', '2')),
    ('Figure 3.1.', pg('fig3_1', '12')),
    ('Figure 3.2.', pg('fig3_bt', '17')),
    ('Figure 3.3.', pg('fig3_2', '18')),
    ('Figure 3.4.', pg('fig3_4', '19')),
    ('Figure 4.1.', pg('fig4_2', '23')),
    ('Figure 4.2.', pg('fig4_faults', '26')),
    ('Figure 5.1.', pg('fig5_1dl', '35')),
    ('Figure 5.2.', pg('fig5_2ht', '36')),
    ('Figure 5.3.', pg('fig5_2sd', '38')),
    ('Figure 5.4.', pg('fig5_6sd', '42')),
]
fill_toc_table(table_after_para('Table of Contents'), toc_entries)
fill_lo_table(table_after_para('List of Tables'), lot_entries)
fill_lo_table(table_after_para('List of Figures'), lof_entries)

doc.save(OUT)
json.dump(keys, open(os.path.join(HERE, 'keys_final.json'), 'w'), ensure_ascii=False, indent=1)
print('saved', OUT)
