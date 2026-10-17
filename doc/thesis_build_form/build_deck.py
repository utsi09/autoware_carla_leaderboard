# -*- coding: utf-8 -*-
"""본심 발표 PPT (13장) — 학과 Presentation Guideline 준수:
Arial bold ≥24pt(쪽번호 제외), 흰 배경·검정 글씨, 5mm 여백, 가로형,
슬라이드당 ≤8줄, 하단 중앙 쪽번호, 로고 없음, 제목 40pt + 회색 괘선."""
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

FB = '/home/taewook/autoware_carla_leaderboard/doc/paper_build/figs'
FR = '/home/taewook/autoware_carla_leaderboard/doc/ral_build/figs'
F = {
 'problem': f'{FB}/fig1_problem.png',   'arch': f'{FR}/fig2b_dualchannel.png',
 'bt': f'{FB}/fig3_bt.png',             'policy': f'{FB}/tab1_policy.png',
 'cand': f'{FB}/fig4_candidates.png',   'scenes': f'{FB}/fig5_scenes.png',
 'faults': f'{FR}/fig6_faults.png',     'latency': f'{FB}/fig7_latency.png',
 'health': f'{FR}/fig12_health.png',    'stopdist': f'{FR}/fig8_stopdist.png',
 'dose': f'{FR}/fig13_speeddose.png',
}
AR = {'problem':1.83,'arch':2.02,'bt':2.23,'policy':3.51,'cand':1.15,'scenes':2.12,
      'faults':1.6,'latency':0.87,'health':1.68,'stopdist':1.26,'dose':2.43}

GREY = RGBColor(0xA6,0xA6,0xA6)
BLACK = RGBColor(0,0,0)
SW, SH = 25.4, 19.05

prs = Presentation()
prs.slide_width = Cm(SW)
prs.slide_height = Cm(SH)
blank = prs.slide_layouts[6]

def txbox(s, x, y, w, h):
    tb = s.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Cm(0.1)
    tf.margin_top = tf.margin_bottom = Cm(0.05)
    return tf

def setrun(r, size, bold=True, italic=False, color=BLACK):
    r.font.name = 'Arial'
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color

def rule(s, y):
    from pptx.enum.shapes import MSO_SHAPE
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(0.5), Cm(y), Cm(SW-1.0), Cm(0.08))
    sh.fill.solid(); sh.fill.fore_color.rgb = GREY
    sh.line.fill.background()
    return sh

def pageno(s, n):
    tf = txbox(s, SW/2-1.5, 18.25, 3.0, 0.7)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    setrun(p.add_run(), 14); p.runs[0].text = str(n)

def header(s, title, n):
    tf = txbox(s, 0.5, 0.45, SW-1.0, 2.0)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title
    setrun(r, 40 if len(title) <= 28 else 34)
    rule(s, 2.62)
    rule(s, 17.95)
    pageno(s, n)

def est_lines(text, size, width_cm):
    cpl = max(10, int(width_cm * 1.93 * 24.0 / size))  # 24pt에서 폭 24.2cm당 ~47자
    import math
    return max(1, math.ceil(len(text) / cpl))

def content(s, items, figkey=None, figw=None, x=0.9):
    w = SW - 2 * x
    y = 3.05
    h = 0.0
    for text, lvl, size in items:
        lines = est_lines(('     ' if lvl else '') + '• ' + text, size, w)
        h += lines * (size * 0.045) + 0.42
    bullets(s, items, x=x, y=y, w=w, h=h + 0.3)
    if figkey:
        top = y + h + 0.35
        avail = 17.75 - top
        fw = figw or (SW - 2)
        fh = fw / AR[figkey]
        if fh > avail:
            fh = avail
            fw = fh * AR[figkey]
        s.shapes.add_picture(F[figkey], Cm((SW - fw) / 2), Cm(top + (avail - fh) / 2),
                             Cm(fw), Cm(fh))

def bullets(s, items, x=0.9, y=3.1, w=SW-1.8, h=14.0):
    tf = txbox(s, x, y, w, h)
    first = True
    for text, lvl, size in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(max(8, size*0.45))
        prefix = '• ' if lvl == 0 else '– '
        r = p.add_run(); r.text = prefix + text
        setrun(r, size)
        # hanging indent (marL/indent, EMU)
        pPr = p._pPr if p._pPr is not None else p.get_or_add_pPr()
        base = 288000 if lvl == 0 else 648000   # 0.8cm / 1.8cm
        pPr.set('marL', str(base))
        pPr.set('indent', str(-288000))
    return tf

def pic(s, key, w, y=None, x=None):
    h = w / AR[key]
    if x is None: x = (SW - w) / 2
    if y is None: y = 17.6 - h
    s.shapes.add_picture(F[key], Cm(x), Cm(y), Cm(w), Cm(h))
    return h

def slide():
    return prs.slides.add_slide(blank)

# ---------- 1. Title ----------
s = slide()
tf = txbox(s, 0.8, 3.3, SW-1.6, 5.0)
for li, line in enumerate(['Continuous Sensor Health Monitoring and',
                           'Risk/Capability-aware Graded MRM Supervision',
                           'for Fail-Safe Operation of Autonomous Vehicles']):
    p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(6)
    r = p.add_run(); r.text = line; setrun(r, 28)
tf = txbox(s, 1.2, 9.6, SW-2.4, 1.4)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = 'Taewook Eum'; setrun(r, 30)
tf = txbox(s, 1.2, 12.2, SW-2.4, 3.9)
for li, (line, size) in enumerate([('Sungkyunkwan University', 26),
                                   ('Department of Electrical and Computer Engineering', 26),
                                   ('Advisor: Prof. Tae Yong Kuc', 24)]):
    p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(10)
    r = p.add_run(); r.text = line; setrun(r, size, italic=True)
pageno(s, 1)

# ---------- 2. Outline ----------
s = slide(); header(s, 'Outline', 2)
bullets(s, [
 ('Background & Problem', 0, 28),
 ('Related Work', 0, 28),
 ('Proposed System: Health Monitoring & MRM Supervision', 0, 28),
 ('Experimental Setup: 1,862-Case Closed Loop', 0, 28),
 ('Results: Detection, Encounters, Causal Cap Effect', 0, 28),
 ('Conclusion', 0, 28),
], y=3.6)

# ---------- 3. Background & Problem ----------
s = slide(); header(s, 'Background & Problem', 3)
content(s, [
 ('Level-4 autonomy must reach a minimal risk condition on its own — ISO 26262 FTTI; SOTIF (ISO 21448)', 0, 24),
 ('Evaluated stock stack: in-lane stop only; faults unmonitored → 238–275 m blind driving', 0, 24),
 ('Design axes: stop now? / how far? / how fast?', 0, 24),
], figkey='problem', figw=13.6)

# ---------- 4. Related Work ----------
s = slide(); header(s, 'Related Work', 4)
bullets(s, [
 ('Adaptive-MRM studies presuppose the fault state (Xue 2018, Tang 2026)', 0, 26),
 ('Selection by traffic TTC/THW', 1, 24),
 ('Validation on representative scenarios', 1, 24),
 ('This work: estimate faults online as a continuous health score', 0, 26),
 ('Graded MRM selection + evacuation speed cap', 1, 24),
 ('1,862-case closed loop with verified collision verdicts', 1, 24),
], y=3.8)

# ---------- 5. Proposed System ----------
s = slide(); header(s, 'Proposed System', 5)
content(s, [
 ('Non-intrusive add-on: read-only monitoring, standard interfaces', 0, 24),
 ('MOVE (reroute + speed cap) vs DEFER (delegate to stock stop)', 0, 24),
], figkey='arch', figw=21.4)

# ---------- 6. Continuous Health Scoring ----------
s = slide(); header(s, 'Continuous Health Scoring', 6)
content(s, [
 ('Freshness · quality · physics scores per sensor at 10 Hz (behavior tree)', 0, 24),
 ('8-sector fusion → health = min(perception, localization)', 0, 24),
 ('Dual thresholds 0.7 / 0.3 → NORMAL / DEGRADED / CRITICAL', 0, 24),
], figkey='bt', figw=19.4)

# ---------- 7. MRM Supervision & Speed Cap ----------
s = slide(); header(s, 'MRM Supervision & Speed Cap', 7)
content(s, [
 ('Refuge by cost: exposure time + traffic-light observability gate', 0, 24),
 ('Speed cap 3.0 m/s from d = vτ + v²/2a  (model bound ≈ 3.9 m/s)', 0, 24),
 ('Hard fails: evidence-aware fast path decides within 1.5 s', 0, 24),
], figkey='policy', figw=24.2)

# ---------- 8. Experimental Setup ----------
s = slide(); header(s, 'Experimental Setup', 8)
content(s, [
 ('CARLA 0.9.16 ↔ Autoware 0.39.0 closed loop (AWCL), Town07, 15 km/h', 0, 24),
 ('1,862 valid cases: fault × severity × traffic × method + encounters', 0, 24),
 ('Full simulator restart every case — independence guaranteed', 0, 24),
], figkey='scenes', figw=18.6)

# ---------- 9. Fault Model & Verdict Reliability ----------
s = slide(); header(s, 'Fault Model & Verdict Reliability', 9)
content(s, [
 ('11 physics-based fault types (LiDAR, camera, GNSS, IMU)', 0, 24),
 ('Pedestrian collision sensor never fires in CARLA', 0, 24),
 ('Dual verdict: ego ledger (RAW-HIT) + physics oracle', 0, 24),
], figkey='faults', figw=17.0)

# ---------- 10. Results: Detection & Graded Response ----------
s = slide(); header(s, 'Results: Detection & Response', 10)
content(s, [
 ('Detection 96–100% (19/20 conditions);  0 false alarms / 45 runs', 0, 24),
 ('Exposure 238–275 m → 11–58 m;  hard-fail decision ≤ 1.20 s', 0, 24),
 ('Health bands separate by severity → tolerate / evacuate / defer', 0, 24),
], figkey='health', figw=14.2)

# ---------- 11. Results: Avoid/Collide Boundary ----------
s = slide(); header(s, 'Results: Avoid/Collide Boundary', 11)
content(s, [
 ('Boundary: 8–11 m at 4.2 m/s (model 8.8 m);  7–8 m at 3.0 m/s', 0, 24),
 ('8 m dart-out: stock 10/10 collide vs settled cruise 8/8 avoid', 0, 24),
 ('Transient encounters 2/2 collide → risk lives in the transient', 0, 24),
], figkey='stopdist', figw=10.6)

# ---------- 12. Results: Causal Cap Effect ----------
s = slide(); header(s, 'Results: Causal Cap Effect', 12)
content(s, [
 ('120 pre-registered runs; settlement trigger — cap holds at encounter', 0, 24),
 ('8 m paired contrast: discordant 10:0, exact McNemar p = 0.002', 0, 24),
 ('Boundary moves inward monotonically: beyond 9 m → about 7 m', 0, 24),
], figkey='dose', figw=20.2)

# ---------- 13. Conclusion ----------
s = slide(); header(s, 'Conclusion', 13)
bullets(s, [
 ('Graded fault response — health score + MRM supervision, 1,862 cases', 0, 25),
 ('Stop vs evacuate: a temporal trade-off; risk lives in the transient', 0, 25),
 ('Settled low speed shifts the boundary (10:0, p = 0.002)', 0, 25),
 ('Reusable verdict methodology: ledger + physics oracle', 0, 25),
 ('Future: high-speed ODD, adaptive baselines, real vehicle', 0, 25),
], y=3.8)

prs.save('defense_deck.pptx')
print('saved defense_deck.pptx, slides:', len(prs.slides.__iter__.__self__._sldIdLst))
