# -*- coding: utf-8 -*-
"""제출 최종본 패처: 사용자 0831 docx → 제출용 정리.
- 유의사항(1쪽)·측면표지 삭제, 6장 절번호 오류(6.→4.), 4장 헤딩 스타일 통일,
  2장 도입문 신구조 반영, TOC/LoT/LoF를 신구조+실측 쪽수로 재작성."""
import os, sys, copy
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import Table

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(HERE, 'user_submit_base.docx')
OUT = '/home/taewook/autoware_carla_leaderboard/doc/thesis_submit_v1_0901.docx'
FONT = 'HY신명조'

doc = Document(BASE)

def find_para(text, start=0, exact=False):
    for i, p in enumerate(doc.paragraphs):
        if i >= start and ((p.text.strip() == text) if exact else (text in p.text)):
            return i, p
    raise KeyError(text)

# ---------- 1) 유의사항 페이지 삭제 ----------
i, p = find_para('학위논문 작성 유의사항')
j, stop = find_para('본 페이지는 전자전기컴퓨터공학과')
el = p._element
while el is not None and el is not stop._element:
    nxt = el.getnext()
    el.getparent().remove(el)
    el = nxt
print('유의사항 페이지 삭제')

# ---------- 2) 측면표지 삭제 (+ 마지막 섹션 정리) ----------
i, sc = find_para('Side Cover')
# 직전의 문단 수준 sectPr(국문요약 섹션 끝)을 찾아 body 최종 sectPr로 승격
body = doc.element.body
prev_sectpr = None
for pp in doc.paragraphs:
    pPr = pp._p.find(qn('w:pPr'))
    if pPr is not None and pPr.find(qn('w:sectPr')) is not None:
        prev_sectpr = (pp, pPr.find(qn('w:sectPr')))
if prev_sectpr is None:
    raise RuntimeError('요약 섹션 sectPr 못 찾음')
last_para_sect, sect_el = prev_sectpr
final_sect = body.find(qn('w:sectPr'))
body.replace(final_sect, copy.deepcopy(sect_el))
pPr = last_para_sect._p.find(qn('w:pPr'))
pPr.remove(sect_el)
# side cover 콘텐츠 삭제: sc 문단부터 body 끝(sectPr 제외)까지
el = sc._element
while el is not None:
    nxt = el.getnext()
    if not el.tag.endswith('}sectPr'):
        el.getparent().remove(el)
    el = nxt
print('측면표지 삭제 + 섹션 정리')

# ---------- 3) 6장 절번호 수정: '6. Limitations' → '4. Limitations' ----------
i, p = find_para('6. Limitations', exact=True)
for t_el in p._p.iter(qn('w:t')):
    if t_el.text and '6. Limitations' in t_el.text:
        t_el.text = t_el.text.replace('6. Limitations', '4. Limitations')
print('6장 절번호 수정')

# ---------- 4) 4장 헤딩 스타일 통일 ----------
for old, new in [('4.1 Experimental Environment', '1. Experimental Environment'),
                 ('4.2 Fault Model', '2. Fault Model')]:
    i, p = find_para(old, exact=True)
    for t_el in p._p.iter(qn('w:t')):
        if t_el.text and old in t_el.text:
            t_el.text = t_el.text.replace(old, new)
print('4장 헤딩 통일')

# ---------- 5) 2장 도입문 ----------
i, p = find_para('six strands (Sections 2.1-2.6)')
for t_el in p._p.iter(qn('w:t')):
    if t_el.text and 'six strands' in t_el.text:
        t_el.text = ('Related work is organized into four strands (Sections 2.1-2.4), '
                     'and Table 2.1 compares them with this study.')
print('2장 도입문 수정')

# ---------- 6) TOC/LoT/LoF 재작성 (실측 쪽수) ----------
TOC_GRID = [2450, 849, 426, 1134, 993, 1276, 809, 567]
LO_GRID = [2127, 5811, 567]
DOTS = {7: 28, 6: 24, 5: 20, 4: 18, 3: 13, 2: 9, 1: 4}

def _mk_tc(width, span, text, jc):
    tc = OxmlElement('w:tc')
    tcPr = OxmlElement('w:tcPr')
    w = OxmlElement('w:tcW'); w.set(qn('w:w'), str(width)); w.set(qn('w:type'), 'dxa'); tcPr.append(w)
    if span > 1:
        gs = OxmlElement('w:gridSpan'); gs.set(qn('w:val'), str(span)); tcPr.append(gs)
    tb = OxmlElement('w:tcBorders')
    for s in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{s}'); e.set(qn('w:val'), 'nil'); tb.append(e)
    tcPr.append(tb)
    tm = OxmlElement('w:tcMar')
    for s, v in (('top', 28), ('left', 102), ('bottom', 28), ('right', 102)):
        e = OxmlElement(f'w:{s}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); tm.append(e)
    tcPr.append(tm)
    va = OxmlElement('w:vAlign'); va.set(qn('w:val'), 'center'); tcPr.append(va)
    tc.append(tcPr)
    pel = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    pPr.append(OxmlElement('w:wordWrap'))
    sg = OxmlElement('w:snapToGrid'); sg.set(qn('w:val'), '0'); pPr.append(sg)
    sp = OxmlElement('w:spacing'); sp.set(qn('w:line'), '348'); sp.set(qn('w:lineRule'), 'auto'); pPr.append(sp)
    if jc:
        j = OxmlElement('w:jc'); j.set(qn('w:val'), jc); pPr.append(j)
    pel.append(pPr)
    rel = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts'); rf.set(qn('w:ascii'), FONT); rf.set(qn('w:eastAsia'), FONT); rPr.append(rf)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '22'); rPr.append(sz)
    rel.append(rPr)
    t_ = OxmlElement('w:t'); t_.set(qn('xml:space'), 'preserve'); t_.text = text; rel.append(t_)
    pel.append(rel)
    tc.append(pel)
    return tc

def _grid(te, grid):
    tblPr = te.find(qn('w:tblPr'))
    for tag in ('w:tblW', 'w:tblBorders', 'w:tblLayout', 'w:tblCellMar'):
        e = tblPr.find(qn(tag))
        if e is not None:
            tblPr.remove(e)
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), str(sum(grid))); w.set(qn('w:type'), 'dxa'); tblPr.append(w)
    bd = OxmlElement('w:tblBorders')
    for s in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{s}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '2')
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), '000000'); bd.append(e)
    for s in ('insideH', 'insideV'):
        e = OxmlElement(f'w:{s}'); e.set(qn('w:val'), 'nil'); bd.append(e)
    tblPr.append(bd)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    og = te.find(qn('w:tblGrid'))
    if og is not None:
        te.remove(og)
    g = OxmlElement('w:tblGrid')
    for cw in grid:
        c = OxmlElement('w:gridCol'); c.set(qn('w:w'), str(cw)); g.append(c)
    te.insert(list(te).index(tblPr) + 1, g)

def fill_toc(tbl, entries):
    te = tbl._tbl
    _grid(te, TOC_GRID)
    for tr in list(te.findall(qn('w:tr'))):
        te.remove(tr)
    cum = [sum(TOC_GRID[:k]) for k in range(1, 8)]
    for text, page in entries:
        est = int(len(text) * 128) + 204
        k = next((i + 1 for i, c in enumerate(cum[:6]) if c >= est), 6)
        tr = OxmlElement('w:tr')
        tr.append(_mk_tc(cum[k - 1], k, text, 'left'))
        tr.append(_mk_tc(sum(TOC_GRID[k:7]), 7 - k, '·' * DOTS[7 - k], 'center'))
        tr.append(_mk_tc(567, 1, str(page), 'center'))
        te.append(tr)

def fill_lo(tbl, entries):
    te = tbl._tbl
    _grid(te, LO_GRID)
    for tr in list(te.findall(qn('w:tr'))):
        te.remove(tr)
    for label, page in entries:
        tr = OxmlElement('w:tr')
        tr.append(_mk_tc(LO_GRID[0], 1, label, 'left'))
        tr.append(_mk_tc(LO_GRID[1], 1, '·' * 25, 'center'))
        tr.append(_mk_tc(567, 1, str(page), 'center'))
        te.append(tr)

def table_after(text):
    i, p = find_para(text)
    el = p._element.getnext()
    while el is not None and not el.tag.endswith('}tbl'):
        el = el.getnext()
    return Table(el, doc)

toc = [
 ('Abstract', 'v'),
 ('1. Introduction', '1'),
 ('1.1  Research Background', '1'),
 ('1.2  Research Questions', '3'),
 ('1.3  Contribution', '4'),
 ('2. Background and Related Work', '5'),
 ('2.1  Standards and Regulatory Frameworks', '6'),
 ('2.2  FTTI and the Derivation of Time Budgets', '6'),
 ('2.3  Fault Detection Theory and Sensor Degradation Models', '7'),
 ('2.4  Fallback and MRM Planning', '7'),
 ('3. Proposed System', '11'),
 ('3.1  System Architecture', '11'),
 ('3.2  Sensor Scoring', '13'),
 ('3.3  Spatial Sectors and Capability Synthesis', '14'),
 ('3.4  Risk-Coupled Dynamic Detection Budget', '15'),
 ('3.5  MRM Supervisor', '17'),
 ('3.6  Conservative Low-Speed MRM Movement (Speed Cap)', '21'),
 ('3.7  Implementation Boundary and Map Annotation', '22'),
 ('4. Experimental Environment and Design', '22'),
 ('4.1  Experimental Environment', '23'),
 ('4.2  Fault Model', '24'),
 ('4.3  Main Campaign Matrix', '27'),
 ('4.4  Pedestrian Dart-out Encounter Injector', '28'),
 ('4.5  Collision Adjudication Methodology', '28'),
 ('4.6  Distance Ladder and 8 m Encounter Experiment Design', '29'),
 ('4.7  Metrics and Statistics', '30'),
 ('4.8  Reproducibility: Key Experimental Parameters', '31'),
 ('5. Experimental Results', '32'),
 ('5.1  Main Campaign: Detection Performance and Per-Fault Behavior', '32'),
 ('5.2  Exposure Metric: Degraded-State Driving Distance', '38'),
 ('5.3  Distance Ladder: Measuring the Avoid/Collide Boundary', '38'),
 ('5.4  The 8 m Encounter: Low-Speed MRM versus Stock', '40'),
 ('5.5  Low-Speed Ladder: Shift of the Boundary', '42'),
 ('5.6  Causal Verification of the Speed Cap: Dose-Response Experiment', '42'),
 ('5.7  Hard-Fail Decision Latency: Verifying the 1.5 s Deadline', '45'),
 ('6. Limitations and Discussion', '46'),
 ('6.1  Temporal Trade-off, Transient Risk, and Speed Cap', '46'),
 ('6.2  Lessons for Adjudication Methodology', '47'),
 ('6.3  The Cost of Design Conservatism and the Execution/Map Layers', '47'),
 ('6.4  Limitations', '48'),
 ('7. Conclusion', '49'),
 ('References', '50'),
 ('논문요약 (Korean Abstract)', '57'),
]
lot = [('Table 2.1.', '9'), ('Table 4.1.', '25'), ('Table 4.2.', '27'), ('Table 4.3.', '31'),
       ('Table 5.1.', '33'), ('Table 5.2.', '37'), ('Table 5.3.', '39'), ('Table 5.4.', '43'), ('Table 5.5.', '45')]
lof = [('Fig. 1.1.', '3'), ('Fig. 3.1.', '12'), ('Fig. 3.2.', '15'), ('Fig. 3.3.', '18'), ('Fig. 3.4.', '20'),
       ('Fig. 4.1.', '23'), ('Fig. 4.2.', '26'), ('Fig. 5.1.', '35'), ('Fig. 5.2.', '37'), ('Fig. 5.3.', '41'), ('Fig. 5.4.', '44')]
fill_toc(table_after('Table of Contents'), toc)
fill_lo(table_after('List of Tables'), lot)
fill_lo(table_after('List of Figures'), lof)
print('TOC/LoT/LoF 재작성')

doc.save(OUT)
print('saved', OUT)
