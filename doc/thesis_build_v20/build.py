# -*- coding: utf-8 -*-
import json, os, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from content1 import (TITLE_KO_L1, TITLE_KO_L2, TITLE_KO_L3,
                      TITLE_EN_L1, TITLE_EN_L2, TITLE_EN_L3,
                      SUMMARY_KO, SUMMARY_KEYWORDS, CH1, CH2)
from content2 import CH3, CH4
from content3 import CH5, CH6, CH7, REFS, ABSTRACT_EN, ABSTRACT_KEYWORDS

TEMPLATE = '/home/taewook/.claude/uploads/3a654bb5-ccec-4bf3-a05d-7c11727be86b/1750d3b5-___011.____________Word_Masters_Template.docx'
OUT = '/home/taewook/autoware_carla_leaderboard/doc/thesis_v21_0827.docx'
HERE = os.path.dirname(os.path.abspath(__file__))
FONT = 'HY신명조'

pagemap = {}
pm_path = os.path.join(HERE, 'pagemap.json')
if os.path.exists(pm_path):
    with open(pm_path) as f:
        pagemap = json.load(f)

keys = {}  # key -> search needle (for page mapping)

# ---------- helpers ----------

def style_run(r, size=11, bold=False):
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    rPr = r._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rPr.insert(0, rf)
    rf.set(qn('w:eastAsia'), FONT)
    return r

GRID = Pt(25.05)  # 서식 템플릿 docGrid linePitch=501 (신명조 11pt 200% 상당)

def para(doc, text='', size=11, bold=False, align=None, spacing=2.0,
         page_break=False, style=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = GRID if spacing == 2.0 else spacing
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if align is not None:
        p.alignment = align
    if page_break:
        pf.page_break_before = True
    if text:
        style_run(p.add_run(text), size=size, bold=bold)
    return p

def body_para(doc, text):
    return para(doc, '   ' + text, size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

def chapter(doc, title, key, first=False):
    keys[key] = title
    para(doc, title, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         page_break=not first)
    para(doc)  # one blank line after big title

def section(doc, title, key=None):
    if key:
        keys[key] = title
    para(doc)  # blank line before mid title
    para(doc, title, size=14, bold=True)
    para(doc)

def subsection(doc, title):
    para(doc, title, size=11, bold=True)

def equation(doc, text, num):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = GRID
    p.paragraph_format.tab_stops.add_tab_stop(Cm(15.0), WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.left_indent = Cm(1.0)
    style_run(p.add_run(text), size=11)
    style_run(p.add_run('\t' + num), size=11)

# 한 쪽을 초과하는 큰 표는 강제 새 쪽 시작 대신 자연 분할을 허용(머리행 반복)
def caption_para(doc, caption):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.3
    pf.space_before = Pt(2); pf.space_after = Pt(0)
    head, sep, rest = caption.partition('. ')
    style_run(p.add_run(head + sep), size=9.5, bold=True)
    if rest:
        style_run(p.add_run(rest), size=9.5)
    return p

ALLOW_SPLIT = {'tbl2_1', 'tbl5_3b', 'tbl5_0'}

def add_table(doc, caption, key, rows, note):
    keys[key] = caption.split('. ')[0]
    cap = caption_para(doc, caption)
    cap.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = 1  # center
    keep_whole = (len(rows) <= 15) and key not in ALLOW_SPLIT
    for ri, tr in enumerate(t.rows):
        trPr = tr._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        if ri == 0:
            trPr.append(OxmlElement('w:tblHeader'))  # 분할 시 머리행 반복
    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            cell = t.cell(ri, ci)
            cell.text = ''
            first = True
            for line in str(cell_text).split('\n'):
                cp = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                cp.paragraph_format.line_spacing = 1.2
                cp.paragraph_format.space_before = Pt(1)
                cp.paragraph_format.space_after = Pt(1)
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
                if keep_whole and ri < len(rows) - 1:
                    cp.paragraph_format.keep_with_next = True
                style_run(cp.add_run(line), size=9.5, bold=(ri == 0))
    if note:
        n = para(doc, note, size=9.5)
        n.paragraph_format.line_spacing = 1.3
    para(doc)

def drop_last_blank(doc):
    # 장 마지막 블록이 표/그림일 때 붙는 말미 공백 문단 제거(공백 쪽 방지)
    last = doc.paragraphs[-1]
    if not last.text.strip():
        last._element.getparent().remove(last._element)

def add_figure(doc, caption, key, path, note, width_cm=None):
    # fig 블록 6번째 요소로 그림별 폭(cm) 오버라이드 지원 (기본 14.5 cm)
    keys[key] = caption.split('. ')[0]
    if path and os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_with_next = True  # 그림-캡션 분리 방지
        r = p.add_run()
        r.add_picture(path, width=Cm(width_cm if width_cm else 14.5))
    cap = caption_para(doc, caption)
    cap.paragraph_format.keep_together = True  # 캡션 쪽 넘김 분할 방지
    if note:
        n = para(doc, note, size=9.5)
        n.paragraph_format.line_spacing = 1.3
    para(doc)

def render_blocks(doc, blocks, first_chapter=False):
    first = first_chapter
    for b in blocks:
        kind = b[0]
        if kind == 'ch':
            chapter(doc, b[1], b[2], first=False)
            first = False
        elif kind == 'sec':
            section(doc, b[1], b[2])
        elif kind == 'sub':
            subsection(doc, b[1])
        elif kind == 'p':
            body_para(doc, b[1])
        elif kind == 'pn':
            para(doc, b[1], size=11)
        elif kind == 'eq':
            equation(doc, b[1], b[2])
        elif kind == 'tbl':
            add_table(doc, b[1], b[2], b[3], b[4])
        elif kind == 'fig':
            add_figure(doc, b[1], b[2], b[3], b[4], b[5] if len(b) > 5 else None)
        elif kind == 'blank':
            para(doc)

def set_pgnum(sec, fmt=None, start=None):
    sectPr = sec._sectPr
    pg = sectPr.find(qn('w:pgNumType'))
    if pg is None:
        pg = OxmlElement('w:pgNumType')
        cols = sectPr.find(qn('w:cols'))
        if cols is not None:
            cols.addprevious(pg)
        else:
            sectPr.append(pg)
    if fmt:
        pg.set(qn('w:fmt'), fmt)
    if start is not None:
        pg.set(qn('w:start'), str(start))

def footer_pagenum(sec, size=11):
    sec.footer.is_linked_to_previous = False
    ftr = sec.footer
    p = ftr.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    r1 = p.add_run(); style_run(r1, size=size)
    fld = OxmlElement('w:fldChar'); fld.set(qn('w:fldCharType'), 'begin')
    r1._element.append(fld)
    r2 = p.add_run(); style_run(r2, size=size)
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    r2._element.append(instr)
    r3 = p.add_run(); style_run(r3, size=size)
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'end')
    r3._element.append(fld2)

def empty_footer(sec):
    sec.footer.is_linked_to_previous = False
    p = sec.footer.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)

def toc_line(doc, text, page_label, indent=0, size=11):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = GRID
    pf.left_indent = Cm(indent)
    pf.tab_stops.add_tab_stop(Cm(15.0 - indent), WD_TAB_ALIGNMENT.RIGHT,
                              WD_TAB_LEADER.DOTS)
    style_run(p.add_run(text), size=size)
    style_run(p.add_run('\t' + (page_label or '')), size=size)

# ---------- document ----------

doc = Document(TEMPLATE)
body = doc.element.body
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)

sec0 = doc.sections[0]
# 잔존 sectPr는 템플릿 마지막 섹션(표지측면, 여백 1cm)이므로
# 본문 섹션(템플릿 1번 섹션) 규격으로 복원: 위/아래 5.3cm, 좌/우 3cm
from docx.shared import Twips
sec0.top_margin = Twips(3005)
sec0.bottom_margin = Twips(3005)
sec0.left_margin = Twips(1701)
sec0.right_margin = Twips(1701)
sec0.header_distance = Twips(851)
sec0.footer_distance = Twips(851)
empty_footer(sec0)

CENTER = WD_ALIGN_PARAGRAPH.CENTER
RIGHT = WD_ALIGN_PARAGRAPH.RIGHT
LEFT = WD_ALIGN_PARAGRAPH.LEFT

def cover_page(doc, english=False, submit=False):
    para(doc, '碩士學位 請求論文', size=16, bold=True, align=LEFT, spacing=1.5)
    para(doc, '指導敎授  ○ ○ ○', size=14, align=RIGHT, spacing=1.5)
    for _ in range(2 if submit else 3):
        para(doc, spacing=1.5)
    para(doc, TITLE_KO_L1, size=22, bold=True, align=CENTER, spacing=1.5)
    para(doc, TITLE_KO_L2, size=22, bold=True, align=CENTER, spacing=1.5)
    para(doc, TITLE_KO_L3, size=22, bold=True, align=CENTER, spacing=1.5)
    if english:
        para(doc, spacing=1.5)
        para(doc, TITLE_EN_L1, size=16, align=CENTER, spacing=1.5)
        para(doc, TITLE_EN_L2, size=16, align=CENTER, spacing=1.5)
        para(doc, TITLE_EN_L3, size=16, align=CENTER, spacing=1.5)
    if submit:
        para(doc, spacing=1.5)
        para(doc, '이 論文을 工學 碩士學位請求論文으로 提出합니다.', size=16,
             align=CENTER, spacing=1.5)
        para(doc, spacing=1.5)
        para(doc, '2026 年 4 月    日', size=16, align=CENTER, spacing=1.5)
    n_blank = 1 if (english and submit) else (5 if english else 7)
    for _ in range(n_blank):
        para(doc, spacing=1.5)
    para(doc, '成均館大學校 一般大學院', size=16, bold=True, align=CENTER, spacing=1.5)
    para(doc, '電子電氣컴퓨터工學科', size=14, align=CENTER, spacing=1.5)
    para(doc, '○ ○ ○', size=16, bold=True, align=CENTER, spacing=1.5)

# 1) 표지
cover_page(doc, english=False, submit=False)

# 2) 속표지
sec = doc.add_section(WD_SECTION.NEW_PAGE)
cover_page(doc, english=True, submit=False)

# 3) 심사청구서
sec = doc.add_section(WD_SECTION.NEW_PAGE)
cover_page(doc, english=True, submit=True)

# 4) 인준지
sec = doc.add_section(WD_SECTION.NEW_PAGE)
for _ in range(2):
    para(doc, spacing=1.5)
para(doc, '이 論文을 ○○○의 工學', size=16, bold=True, align=CENTER)
para(doc, '碩士學位 論文으로 認定함.', size=16, bold=True, align=CENTER)
for _ in range(2):
    para(doc, spacing=1.5)
para(doc, '2026 年 6 月    日', size=16, align=CENTER)
for _ in range(3):
    para(doc, spacing=1.5)
for role in ['審査委員長', '審査委員    ', '審査委員    ']:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = GRID
    p.alignment = RIGHT
    style_run(p.add_run(role + '                              (印)'), size=14)
    para(doc, spacing=1.5)

# 5) 목차 ~ 논문요약 (로마 숫자)
sec = doc.add_section(WD_SECTION.NEW_PAGE)
set_pgnum(sec, fmt='lowerRoman', start=1)
footer_pagenum(sec)

para(doc, '목        차', size=16, bold=True, align=CENTER)
para(doc)

ALL_BLOCKS = CH1 + CH2 + CH3 + CH4 + CH5 + CH6 + CH7
for b in ALL_BLOCKS:
    if b[0] == 'ch':
        toc_line(doc, b[1].replace('  ', ' '), pagemap.get(b[2], ''))
    elif b[0] == 'sec':
        toc_line(doc, b[1], pagemap.get(b[2], ''), indent=1.0)
toc_line(doc, '참 고 문 헌', pagemap.get('refs', ''))
toc_line(doc, 'ABSTRACT', pagemap.get('abstract', ''))

# 표목차
para(doc, '표 목 차', size=14, bold=True, align=CENTER, page_break=True)
para(doc)
for b in ALL_BLOCKS:
    if b[0] == 'tbl':
        cap = b[1]
        parts = cap.split('. ', 1)
        title = parts[0] + '. ' + parts[1].split('. ')[0] if len(parts) > 1 else cap
        toc_line(doc, title, pagemap.get(b[2], ''))

# 그림목차
para(doc, '그 림 목 차', size=14, bold=True, align=CENTER, page_break=True)
para(doc)
for b in ALL_BLOCKS:
    if b[0] == 'fig':
        cap = b[1]
        parts = cap.split('. ', 1)
        title = parts[0] + '. ' + parts[1].split('. ')[0] if len(parts) > 1 else cap
        toc_line(doc, title, pagemap.get(b[2], ''))

# 논문요약
para(doc, '논문요약', size=14, bold=True, page_break=True)
para(doc)
para(doc, TITLE_KO_L1 + ' ' + TITLE_KO_L2, size=16, bold=True, align=CENTER, spacing=1.5)
para(doc, TITLE_KO_L3, size=16, bold=True, align=CENTER, spacing=1.5)
para(doc)
for t in SUMMARY_KO:
    body_para(doc, t)
para(doc)
para(doc, SUMMARY_KEYWORDS, size=11, bold=True)

# 6) 본문 (아라비아 숫자, 1부터)
sec = doc.add_section(WD_SECTION.NEW_PAGE)
set_pgnum(sec, fmt='decimal', start=1)
footer_pagenum(sec)

first = True
for chap in [CH1, CH2, CH3, CH4, CH5, CH6, CH7]:
    for i, b in enumerate(chap):
        if b[0] == 'ch':
            keys[b[2]] = b[1]
            para(doc, b[1], size=16, bold=True, align=CENTER, page_break=not first)
            para(doc)
            first = False
        elif b[0] == 'sec':
            section(doc, b[1], b[2])
        elif b[0] == 'sub':
            subsection(doc, b[1])
        elif b[0] == 'p':
            body_para(doc, b[1])
        elif b[0] == 'eq':
            equation(doc, b[1], b[2])
        elif b[0] == 'tbl':
            add_table(doc, b[1], b[2], b[3], b[4])
            if i == len(chap) - 1:
                drop_last_blank(doc)
        elif b[0] == 'fig':
            add_figure(doc, b[1], b[2], b[3], b[4], b[5] if len(b) > 5 else None)
            if i == len(chap) - 1:
                drop_last_blank(doc)

# 참고문헌
keys['refs'] = '참 고 문 헌'
para(doc, '참 고 문 헌', size=16, bold=True, align=CENTER, page_break=True)
para(doc)
for i, ref in enumerate(REFS, 1):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.left_indent = Cm(0.9)
    pf.first_line_indent = Cm(-0.9)
    pf.space_after = Pt(4)
    style_run(p.add_run(f'[{i}] {ref}'), size=11)

# ABSTRACT
keys['abstract'] = 'ABSTRACT'
para(doc, 'ABSTRACT', size=14, bold=True, page_break=True)
para(doc)
para(doc, TITLE_EN_L1, size=16, bold=True, align=CENTER, spacing=1.5)
para(doc, TITLE_EN_L2, size=16, bold=True, align=CENTER, spacing=1.5)
para(doc, TITLE_EN_L3, size=16, bold=True, align=CENTER, spacing=1.5)
para(doc)
para(doc, '○ ○ ○', size=11, align=RIGHT, spacing=1.5)
para(doc, 'Department of Electrical and Computer Engineering', size=11, align=RIGHT, spacing=1.5)
para(doc, 'The Graduate School', size=11, align=RIGHT, spacing=1.5)
para(doc, 'Sungkyunkwan University', size=11, align=RIGHT, spacing=1.5)
para(doc)
for t in ABSTRACT_EN:
    body_para(doc, '  ' + t.strip())
para(doc)
para(doc, ABSTRACT_KEYWORDS, size=11, bold=True)

doc.save(OUT)
with open(os.path.join(HERE, 'keys.json'), 'w') as f:
    json.dump(keys, f, ensure_ascii=False, indent=1)
print('saved', OUT)
